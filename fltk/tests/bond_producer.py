#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""Independent ODS/XLSX bond valuations, duration caches and partial-status preservation."""
import argparse,datetime as dt,json,math,os,re,subprocess,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableCell,TableRow
from odf.text import P
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from bond_reference import bond_values
ROOT=Path(__file__).resolve().parents[2]
DATES=[dt.date(2008,2,15),dt.date(2017,11,15),dt.date(2007,11,15)]
FORMULAS=['PRICE([.A1];[.A2];.0575;[.B1];100;2;1)',
          'YIELD([.A1];[.A2];.0575;[.C1];100;2;1)',
          'DURATION([.A1];[.A2];.0575;[.B1];2;1)',
          'MDURATION([.A1];[.A2];.0575;[.B1];2;1)',
          'DURATION_ADD([.A1];[.A2];.0575;[.B1];2;1)',
          'DURATION(.1;1000;2000)',
          'PRICE([.A1];[.A2];0;[.B1];100;2;1)',
          'YIELD([.A1];[.A2];0;[.I1];100;2;1)',
          'IFERROR(PRICE([.A1];[.A2];.05;-1;100;2);77)',
          'UNKNOWN_BOND_FUNCTION()',
          'ODDLPRICE([.A1];[.A2];[.A3];.05;.07;100;2)','IFERROR([.M1];77)']
def expected(rate=.065):
    price,duration,modified=bond_values(DATES[0],DATES[1],.0575,rate,100,2,1)
    zero,_,_=bond_values(DATES[0],DATES[1],0,rate,100,2,1)
    return [DATES[0].isoformat(),rate,float(price),rate,float(duration),float(modified),
            float(duration),math.log(2)/math.log1p(.1),
            float(zero),rate,77,'#NAME?','#NAME?',77]
def produce(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=OpenDocumentSpreadsheet();table=Table(name='Bonds');doc.spreadsheet.addElement(table)
    for r,date in enumerate(DATES):
        row=TableRow();table.addElement(row);row.addElement(TableCell(valuetype='date',datevalue=date.isoformat()))
        row.addElement(TableCell(valuetype='float',value=.065 if r==0 else 0))
        if r==0:
            for formula in FORMULAS:
                cell=TableCell(formula='of:='+formula,valuetype='float',value=901);cell.addElement(P(text='901'));row.addElement(cell)
    ods=directory/'bond-valuations.ods';doc.save(str(ods));paths=[ods]
    for mac in (False,True):
        book=Workbook();ws=book.active;ws.title='Bonds'
        if mac:book.epoch=CALENDAR_MAC_1904
        for r,date in enumerate(DATES,1):
            ws.cell(r,1,date);ws.cell(r,1).number_format='yyyy-mm-dd';ws.cell(r,2,.065 if r==1 else 0)
        for c,formula in enumerate(FORMULAS,3):
            f=re.sub(r'\[\.([A-Z]+[0-9]+)\]',lambda m:m[1],formula).replace(';',',')
            if c==7:f=f.replace('DURATION_ADD','DURATION')
            if c==8:f='_xlfn.PDURATION(.1,1000,2000)'
            ws.cell(1,c,'='+f)
        path=directory/('bonds-1904.xlsx' if mac else 'bonds-1900.xlsx');book.save(path)
        with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}';xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for cell in xml.iter(ns+'c'):
            if cell.find(ns+'f') is not None:
                v=cell.find(ns+'v')
                if v is None:v=ET.SubElement(cell,ns+'v')
                v.text='901'
        parts['xl/worksheets/sheet1.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,value in parts.items():z.writestr(name,value)
        paths.append(path)
    return paths
def verify(path,test_bin_dir,directory,rate=.065):
    path=Path(path);directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    before=path.read_bytes();saved=directory/(path.stem+'-recalculated.ods')
    out=subprocess.check_output([str(Path(test_bin_dir).resolve()/'calligra-native-financial-test'),'--document',str(path),str(saved)],text=True,
                                env={**os.environ,'LD_LIBRARY_PATH':'/usr/local/lib/domed/bitmap'})
    rows=json.loads(out);incomplete={12,13,14}
    for column,(got,want) in enumerate(zip(rows[0],expected(rate)),1):
        if column==1:assert got['display'] in (want,want+'T00:00:00.000Z'),(path,column,got,want)
        elif isinstance(want,str):assert got['display']==want,(path,column,got,want)
        else:assert got['kind']==1 and math.isclose(got['value'],want,rel_tol=3e-12,abs_tol=1e-10),(path,column,got,want)
        assert got['incomplete']==(column in incomplete),(column,got)
    for row,date in zip(rows,DATES):assert row[0]['value']==(date-dt.date(1899,12,30)).days
    office='{urn:oasis:names:tc:opendocument:xmlns:office:1.0}';table='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    with zipfile.ZipFile(saved) as z:
        root=ET.fromstring(z.read('content.xml'));cells=next(root.iter(table+'table-row')).findall(table+'table-cell')
        for c in incomplete:assert cells[c-1].get(office+'value')=='901',(c,cells[c-1].attrib)
        assert cells[0].get(office+'date-value') in (DATES[0].isoformat(),DATES[0].isoformat()+'T00:00:00.000Z')
        for c,want in enumerate(expected(rate),1):
            if isinstance(want,str) or c in incomplete:continue
            assert math.isclose(float(cells[c-1].get(office+'value')),want,rel_tol=3e-12,abs_tol=1e-10),(c,cells[c-1].attrib)
    assert path.read_bytes()==before
    (directory/(path.stem+'-report.json')).write_text(json.dumps(rows,indent=2)+'\n')
    print('PASS',path.name,'native bond values, complete DURATION caches, date epochs and retained unsupported caches',flush=True)
    return expected(rate)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('directory',type=Path);p.add_argument('--test-bin-dir',type=Path,default=ROOT/'build-native/fltk');a=p.parse_args()
    for path in produce(a.directory):verify(path,a.test_bin_dir,a.directory)
