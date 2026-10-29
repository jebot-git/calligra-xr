#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""Dated ODS/XLSX cash flows, both Excel epochs and retained package caches."""
import argparse,datetime as dt,json,math,os,re,subprocess,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from scipy.optimize import brentq
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableCell,TableRow
from odf.text import P
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
ROOT=Path(__file__).resolve().parents[2]
DATES=[dt.date(2024,1,1),dt.date(2024,7,1),dt.date(2024,12,31)]
FORMULAS=['XIRR([.B1:.B3];[.A1:.A3])','XNPV(.1;[.B1:.B3];[.A1:.A3])','[.C1]+.01',
          'XNPV([.C1];[.B1:.B3];[.A1:.A3])','IFERROR(XIRR({1;2};{0;365});77)',
          'UNKNOWN_DATED_FUNCTION()','DURATION(DATE(2026;1;1);DATE(2030;1;1);.05;.07;2)',
          'IFERROR([.I1];77)','XIRR({-100;110};{DATE(2024;1;1);DATE(2024;12;31)})']
def expected(cash=-10000):
    def npv(r):return cash+6000/(1+r)**(182/365)+7000/(1+r)
    rate=brentq(npv,-.999999,10,xtol=1e-13,rtol=1e-14)
    return [DATES[0].isoformat(),cash,rate,npv(.1),rate+.01,0,77,'#NAME?','#NAME?',77,.1]
def produce(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=OpenDocumentSpreadsheet();table=Table(name='Dated');doc.spreadsheet.addElement(table)
    for r,(date,cash) in enumerate(zip(DATES,[-10000,6000,7000])):
        row=TableRow();table.addElement(row)
        row.addElement(TableCell(valuetype='date',datevalue=date.isoformat()))
        row.addElement(TableCell(valuetype='float',value=cash))
        if r==0:
            for formula in FORMULAS:
                cell=TableCell(formula='of:='+formula,valuetype='float',value=901);cell.addElement(P(text='901'));row.addElement(cell)
    ods=directory/'dated-cashflows.ods';doc.save(str(ods));paths=[ods]
    for mac in (False,True):
        book=Workbook();ws=book.active;ws.title='Dated'
        if mac:book.epoch=CALENDAR_MAC_1904
        for r,(date,cash) in enumerate(zip(DATES,[-10000,6000,7000]),1):
            ws.cell(r,1,date);ws.cell(r,1).number_format='yyyy-mm-dd';ws.cell(r,2,cash)
        for c,formula in enumerate(FORMULAS,3):
            f=re.sub(r'\[\.([A-Z]+[0-9]+)(?::\.([A-Z]+[0-9]+))?\]',lambda m:m[1]+(':'+m[2] if m[2] else ''),formula).replace(';',',')
            # Excel constants contain literals; CHOOSE constructs the same
            # date array without the retained Calligra expression extension.
            if c==11:f='XIRR({-100,110},CHOOSE({1,2},A1,A3))'
            ws.cell(1,c,'='+f)
        path=directory/('dated-1904.xlsx' if mac else 'dated-1900.xlsx');book.save(path)
        with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for cell in xml.iter(ns+'c'):
            if cell.find(ns+'f') is not None:
                value=cell.find(ns+'v')
                if value is None:value=ET.SubElement(cell,ns+'v')
                value.text='901'
        parts['xl/worksheets/sheet1.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,value in parts.items():z.writestr(name,value)
        paths.append(path)
    return paths
def verify(path,test_bin_dir,directory,cash=-10000):
    path=Path(path);directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    before=path.read_bytes();saved=directory/(path.stem+'-recalculated.ods')
    out=subprocess.check_output([str(Path(test_bin_dir).resolve()/'calligra-native-financial-test'),'--document',str(path),str(saved)],text=True,
                                env={**os.environ,'LD_LIBRARY_PATH':'/usr/local/lib/domed/bitmap'})
    rows=json.loads(out)
    for column,(got,want) in enumerate(zip(rows[0],expected(cash)),1):
        if column==1:
            assert got['display'] in (want,want+'T00:00:00.000Z'),(path,column,got,want)
        elif isinstance(want,str):assert got['display']==want,(path,column,got,want)
        else:assert got['kind']==1 and math.isclose(got['value'],want,rel_tol=3e-10,abs_tol=1e-8),(path,column,got,want)
        assert got['incomplete']==(column in (8,9,10)),(column,got)
    for row,date in zip(rows,DATES):assert row[0]['value']==(date-dt.date(1899,12,30)).days,(path,row[0])
    office='{urn:oasis:names:tc:opendocument:xmlns:office:1.0}';table='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    with zipfile.ZipFile(saved) as z:
        root=ET.fromstring(z.read('content.xml'));cells=next(root.iter(table+'table-row')).findall(table+'table-cell')
        assert cells[0].get(office+'date-value') in ('2024-01-01','2024-01-01T00:00:00.000Z'),cells[0].attrib
        for c in (8,9,10):assert cells[c-1].get(office+'value')=='901',(c,cells[c-1].attrib)
        for c,want in enumerate(expected(cash),1):
            if isinstance(want,str) or c in (8,9,10):continue
            assert math.isclose(float(cells[c-1].get(office+'value')),want,rel_tol=3e-10,abs_tol=1e-8),(c,cells[c-1].attrib)
    assert path.read_bytes()==before
    (directory/(path.stem+'-report.json')).write_text(json.dumps(rows,indent=2)+'\n')
    print('PASS',path.name,'dated cash-flow recalculation, epoch conversion, DATE array expressions and preserved unsupported caches',flush=True)
    return expected(cash)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('directory',type=Path)
    p.add_argument('--test-bin-dir',type=Path,default=ROOT/'build-native/fltk');a=p.parse_args()
    for path in produce(a.directory):verify(path,a.test_bin_dir,a.directory)
