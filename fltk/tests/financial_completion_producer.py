#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""Independent ODS/XLSX producers covering the final retained financial entries."""
import argparse,datetime as dt,json,math,os,re,subprocess,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableCell,TableRow
from odf.text import P
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from financial_completion_reference import accrue,amort,maturity,treasury,euro
ROOT=Path(__file__).resolve().parents[2]
DATES=[dt.date(2004,2,1),dt.date(2004,4,1),dt.date(2004,5,1),dt.date(2004,12,31)]
FORMULAS=['ACCRINT([.B1];[.C1];[.D1];.1;[.A1];4;0)',
          'AMORLINC([.A1];[.B1];[.E1];.1*[.A1];0;.1;1)',
          'AMORDEGRC([.A1];[.B1];[.E1];.1*[.A1];1;.15;1)',
          'EURO("ATS")','EUROCONVERT([.A1];"EUR";"DEM")',
          'ORG.OPENOFFICE.CONVERT([.A1];"EUR";"ATS")',
          'ODDLPRICE([.D1];[.E1];[.B1];.05;.07;[.A1];2;1)',
          'ODDLYIELD([.D1];[.E1];[.B1];.05;[.L1];[.A1];2;1)',
          'PRICEMAT([.D1];[.E1];[.B1];.05;.07;1)',
          'YIELDMAT([.D1];[.E1];[.B1];.05;[.N1];1)',
          'TBILLEQ([.D1];[.E1];.05)','TBILLPRICE([.D1];[.E1];.05)',
          'TBILLYIELD([.D1];[.E1];[.Q1])',
          'UNKNOWN_NATIVE_FINANCIAL()','IFERROR([.S1];77)',
          'ODDFPRICE(DATE(2026;1;1);DATE(2030;1;1);DATE(2025;7;1);DATE(2026;7;1);.05;.07;100;2)']
def expected(money=1000):
    issue,first,start,end=DATES
    odd=maturity('ODDLPRICE',start,end,issue,.05,.07,money,2,1)
    price=maturity('PRICEMAT',start,end,issue,.05,.07,basis=1)
    tbill=treasury('TBILLPRICE',start,end,.05)
    values=[money,*[d.isoformat() for d in DATES],
            accrue(issue,first,start,.1,money,4,0),
            amort('AMORLINC',money,issue,end,.1*money,0,.1,1),
            amort('AMORDEGRC',money,issue,end,.1*money,1,.15,1),
            13.7603,euro(money,'EUR','DEM'),euro(money,'EUR','ATS',True),
            odd,.07,price,.07,treasury('TBILLEQ',start,end,.05),tbill,
            treasury('TBILLYIELD',start,end,tbill),'#NAME?',77,'#NAME?']
    return [v if isinstance(v,str) else float(v) for v in values]
def produce(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=OpenDocumentSpreadsheet();sheet=Table(name='Financial completion');doc.spreadsheet.addElement(sheet)
    row=TableRow();sheet.addElement(row);row.addElement(TableCell(valuetype='float',value=1000))
    for date in DATES:row.addElement(TableCell(valuetype='date',datevalue=date.isoformat()))
    for formula in FORMULAS:
        cell=TableCell(formula='of:='+formula,valuetype='float',value=901);cell.addElement(P(text='901'));row.addElement(cell)
    ods=directory/'financial-completion.ods';doc.save(str(ods));paths=[ods]
    for mac in (False,True):
        book=Workbook();ws=book.active;ws.title='Financial completion'
        if mac:book.epoch=CALENDAR_MAC_1904
        ws.cell(1,1,1000)
        for column,date in enumerate(DATES,2):ws.cell(1,column,date);ws.cell(1,column).number_format='yyyy-mm-dd'
        for column,formula in enumerate(FORMULAS,6):
            f=re.sub(r'\[\.([A-Z]+[0-9]+)\]',lambda m:m[1],formula).replace(';',',')
            if column==9:f='EUROCONVERT(1,"EUR","ATS",TRUE())'
            if column==11:f='EUROCONVERT(A1,"EUR","ATS",TRUE())'
            ws.cell(1,column,'='+f)
        path=directory/('completion-1904.xlsx' if mac else 'completion-1900.xlsx');book.save(path)
        with zipfile.ZipFile(path) as z:parts={name:z.read(name) for name in z.namelist()}
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}';xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for cell in xml.iter(ns+'c'):
            if cell.find(ns+'f') is not None:
                value=cell.find(ns+'v')
                if value is None:value=ET.SubElement(cell,ns+'v')
                value.text='901'
        parts['xl/worksheets/sheet1.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,value in parts.items():z.writestr(name,value)
        paths.append(path)
    return paths
def verify(path,test_bin_dir,directory,money=1000):
    path=Path(path);directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    before=path.read_bytes();saved=directory/(path.stem+'-recalculated.ods')
    output=subprocess.check_output([str(Path(test_bin_dir).resolve()/'calligra-native-financial-test'),'--document',str(path),str(saved)],text=True,
                                   env={**os.environ,'LD_LIBRARY_PATH':'/usr/local/lib/domed/bitmap'})
    rows=json.loads(output);incomplete={19,20,21};dates={2,3,4,5}
    for column,(got,want) in enumerate(zip(rows[0],expected(money)),1):
        if column in dates:
            assert got['display'] in (want,want+'T00:00:00.000Z'),(path,column,got,want)
            assert got['value']==(DATES[column-2]-(dt.date(1904,1,1) if path.suffix=='.xlsx' and '1904' in path.stem else dt.date(1899,12,30))).days,(path,column,got)
        elif isinstance(want,str):assert got['display']==want,(path,column,got,want)
        else:assert got['kind']==1 and math.isclose(got['value'],want,rel_tol=5e-12,abs_tol=1e-8),(path,column,got,want)
        assert got['incomplete']==(column in incomplete),(column,got)
    office='{urn:oasis:names:tc:opendocument:xmlns:office:1.0}';table='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    with zipfile.ZipFile(saved) as z:
        root=ET.fromstring(z.read('content.xml'));cells=next(root.iter(table+'table-row')).findall(table+'table-cell')
        for c in incomplete:assert cells[c-1].get(office+'value')=='901',(c,cells[c-1].attrib)
        for c,want in enumerate(expected(money),1):
            if c in dates:assert cells[c-1].get(office+'date-value') in (want,want+'T00:00:00.000Z')
            elif c not in incomplete and not isinstance(want,str):assert math.isclose(float(cells[c-1].get(office+'value')),want,rel_tol=5e-12,abs_tol=1e-8),(c,cells[c-1].attrib)
    assert path.read_bytes()==before
    (directory/(path.stem+'-report.json')).write_text(json.dumps(rows,indent=2)+'\n')
    print('PASS',path.name,'all final financial entries, epochs, calculated numeric caches and preserved unknown/ODDF caches',flush=True)
    return expected(money)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('directory',type=Path);p.add_argument('--test-bin-dir',type=Path,default=ROOT/'build-native/fltk');a=p.parse_args()
    for path in produce(a.directory):verify(path,a.test_bin_dir,a.directory)
