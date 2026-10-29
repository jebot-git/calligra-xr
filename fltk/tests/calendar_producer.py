#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""Independent ODS/XLSX calendar producers, recalculation and cache preservation."""
import argparse,datetime as dt,json,math,os,re,subprocess,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableCell,TableRow
from odf.text import P
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
ROOT=Path(__file__).resolve().parents[2]
DATES=[dt.date(2024,2,1),dt.date(2024,5,1),dt.date(2024,2,29)]
FORMULAS=['YEARFRAC([.A1];[.A2];0)','DAYS360([.A1];[.A2])',
          'ACCRINTM([.A1];[.A2];.1;[.B1];0)','DISC([.A1];[.A2];.95*[.B1];[.B1];2)',
          'INTRATE([.A1];[.A2];.95*[.B1];[.B1];2)','PRICEDISC([.A1];[.A2];.1;[.B1];0)',
          'RECEIVED([.A1];[.A2];[.B1];.1;0)','YIELDDISC([.A1];[.A2];.95*[.B1];[.B1];2)',
          'IFERROR(YEARFRAC([.A1];[.A2];5);77)','UNKNOWN_CALENDAR_FUNCTION()',
          'ODDFPRICE(DATE(2026;1;1);DATE(2030;1;1);DATE(2025;7;1);DATE(2026;7;1);.05;.07;100;2)','IFERROR([.M1];77)']
def expected(cash=1000):
    return [DATES[0].isoformat(),cash,.25,90,cash*.025,.2,.05/.95*4,
            cash*.975,cash/.975,.05/.95*4,77,'#NAME?','#NAME?',77]
def produce(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=OpenDocumentSpreadsheet();table=Table(name='Calendar');doc.spreadsheet.addElement(table)
    for r,date in enumerate(DATES):
        row=TableRow();table.addElement(row);row.addElement(TableCell(valuetype='date',datevalue=date.isoformat()))
        row.addElement(TableCell(valuetype='float',value=1000 if r==0 else 0))
        if r==0:
            for formula in FORMULAS:
                cell=TableCell(formula='of:='+formula,valuetype='float',value=901);cell.addElement(P(text='901'));row.addElement(cell)
    ods=directory/'calendar-securities.ods';doc.save(str(ods));paths=[ods]
    for mac in (False,True):
        book=Workbook();ws=book.active;ws.title='Calendar'
        if mac:book.epoch=CALENDAR_MAC_1904
        for r,date in enumerate(DATES,1):
            ws.cell(r,1,date);ws.cell(r,1).number_format='yyyy-mm-dd';ws.cell(r,2,1000 if r==1 else 0)
        for c,formula in enumerate(FORMULAS,3):
            f=re.sub(r'\[\.([A-Z]+[0-9]+)\]',lambda m:m[1],formula).replace(';',',')
            ws.cell(1,c,'='+f)
        path=directory/('calendar-1904.xlsx' if mac else 'calendar-1900.xlsx');book.save(path)
        with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}';xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for cell in xml.iter(ns+'c'):
            if cell.find(ns+'f') is not None:
                v=cell.find(ns+'v')
                if v is None:v=ET.SubElement(cell,ns+'v')
                v.text='901'
        parts['xl/worksheets/sheet1.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,value in parts.items():z.writestr(name,value)
        paths.append(path)
    return paths
def verify(path,test_bin_dir,directory,cash=1000):
    path=Path(path);directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    before=path.read_bytes();saved=directory/(path.stem+'-recalculated.ods')
    out=subprocess.check_output([str(Path(test_bin_dir).resolve()/'calligra-native-financial-test'),'--document',str(path),str(saved)],text=True,
                                env={**os.environ,'LD_LIBRARY_PATH':'/usr/local/lib/domed/bitmap'})
    rows=json.loads(out);incomplete={12,13,14}
    for column,(got,want) in enumerate(zip(rows[0],expected(cash)),1):
        if column==1:assert got['display'] in (want,want+'T00:00:00.000Z'),(path,column,got,want)
        elif isinstance(want,str):assert got['display']==want,(path,column,got,want)
        else:assert got['kind']==1 and math.isclose(got['value'],want,rel_tol=3e-12,abs_tol=1e-10),(path,column,got,want)
        assert got['incomplete']==(column in incomplete),(column,got)
    for row,date in zip(rows,DATES):assert row[0]['value']==(date-dt.date(1899,12,30)).days
    office='{urn:oasis:names:tc:opendocument:xmlns:office:1.0}';table='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    with zipfile.ZipFile(saved) as z:
        root=ET.fromstring(z.read('content.xml'));cells=next(root.iter(table+'table-row')).findall(table+'table-cell')
        for c in incomplete:assert cells[c-1].get(office+'value')=='901',(c,cells[c-1].attrib)
        assert cells[0].get(office+'date-value') in (DATES[0].isoformat(),DATES[0].isoformat()+'T00:00:00.000Z')
        for c,want in enumerate(expected(cash),1):
            if isinstance(want,str) or c in incomplete:continue
            assert math.isclose(float(cells[c-1].get(office+'value')),want,rel_tol=3e-12,abs_tol=1e-10),(c,cells[c-1].attrib)
    assert path.read_bytes()==before
    (directory/(path.stem+'-report.json')).write_text(json.dumps(rows,indent=2)+'\n')
    print('PASS',path.name,'native calendar/securities, date epochs, recalculated caches and preserved unsupported caches',flush=True)
    return expected(cash)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('directory',type=Path);p.add_argument('--test-bin-dir',type=Path,default=ROOT/'build-native/fltk');a=p.parse_args()
    for path in produce(a.directory):verify(path,a.test_bin_dir,a.directory)
