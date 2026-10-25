#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""Independent producers: uv run --with python-docx --with python-pptx python ..."""
from pathlib import Path
import base64
import subprocess
import tempfile
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
ROOT = Path(__file__).resolve().parents[2]
with tempfile.TemporaryDirectory() as directory:
    base = Path(directory)
    docx, pptx = base / 'producer.docx', base / 'producer.pptx'
    document = Document()
    document.add_heading('Český 日本語', 0)
    document.add_paragraph('Ordinary paragraph with ').add_run('bold text').bold = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Table cell'
    table.cell(0, 1).text = 'Other cell'
    document.save(docx)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text = 'Český 日本語'
    image = base / 'pixel.png'
    image.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAEklEQVR4nGNQaHig0PCAAUIBACcOBgEpwnGPAAAAAElFTkSuQmCC'))
    slide.shapes.add_picture(str(image), Inches(2), Inches(3), Inches(1), Inches(1))
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = 'Layout placeholder title'
    jpeg = base / 'pixel.jpg'
    Image.new('RGB', (2, 2), 'blue').save(jpeg)
    slide.shapes.add_picture(str(jpeg), Inches(2), Inches(4), Inches(1), Inches(1))
    presentation.save(pptx)
    before = docx.read_bytes(), pptx.read_bytes()
    subprocess.run([str(ROOT / 'build-fltk/fltk/calligra-fltk-office-test'), '--external', str(docx), str(pptx)], check=True)
    assert before == (docx.read_bytes(), pptx.read_bytes())
