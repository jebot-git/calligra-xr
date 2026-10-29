#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""ODS/XLSX coupon schedules, typed date caches and unsupported-cache retention."""
import argparse,datetime as dt,json,math,os,re,subprocess,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableCell,TableRow
from odf.text import P
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
ROOT=Path(__file__).resolve().parents[2]
DATES=[dt.date(2024,8,30),dt.date(2025,2,28),dt.date(2024,2,29)]
FORMULAS=['COUPDAYBS([.A1];[.A2];[.B1];1)','COUPDAYS([.A1];[.A2];[.B1];1)',
          'COUPDAYSNC([.A1];[.A2];[.B1];1)','COUPNCD([.A1];[.A2];[.B1];1)',
          'COUPNUM([.A1];[.A2];[.B1];1)','COUPPCD([.A1];[.A2];[.B1];1)',
          '[.F1]-[.H1]','COUPDAYSNC([.A1];[.A2];[.B1];0)',
          'COUPDAYSNC([.A1];[.A2];[.B1];5)',
          'IFERROR(COUPNUM([.A1];[.A2];5);77)','UNKNOWN_COUPON_FUNCTION()',
          'ODDFPRICE(DATE(2026;1;1);DATE(2030;1;1);DATE(2025;7;1);DATE(2026;7;1);.05;.07;100;2)','IFERROR([.N1];77)',
          '[.F1]=DATE(2024;8;31)']
def expected(frequency=2):
    assert frequency in (2,4)
    prev=dt.date(2024,2,29) if frequency==2 else dt.date(2024,5,31)
    next_=dt.date(2024,8,31);days=(next_-prev).days
    return [DATES[0].isoformat(),frequency,(DATES[0]-prev).days,days,1,
            next_.isoformat(),2 if frequency==2 else 3,prev.isoformat(),days,0,1,
            77,'#NAME?','#NAME?',77,True]
def produce(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=OpenDocumentSpreadsheet();table=Table(name='Coupons');doc.spreadsheet.addElement(table)
    for r,date in enumerate(DATES):
        row=TableRow();table.addElement(row);row.addElement(TableCell(valuetype='date',datevalue=date.isoformat()))
        row.addElement(TableCell(valuetype='float',value=2 if r==0 else 0))
        if r==0:
            for formula in FORMULAS:
                cell=TableCell(formula='of:='+formula,valuetype='float',value=901);cell.addElement(P(text='901'));row.addElement(cell)
    ods=directory/'coupon-schedules.ods';doc.save(str(ods));paths=[ods]
    for mac in (False,True):
        book=Workbook();ws=book.active;ws.title='Coupons'
        if mac:book.epoch=CALENDAR_MAC_1904
        for r,date in enumerate(DATES,1):
            ws.cell(r,1,date);ws.cell(r,1).number_format='yyyy-mm-dd';ws.cell(r,2,2 if r==1 else 0)
        for c,formula in enumerate(FORMULAS,3):
            f=re.sub(r'\[\.([A-Z]+[0-9]+)\]',lambda m:m[1],formula).replace(';',',')
            # Basis 5 is a Calligra extension; the standard Excel document uses
            # the equivalent one-day actual convention for this boundary case.
            if c==11:f='COUPDAYSNC(A1,A2,B1,1)'
            ws.cell(1,c,'='+f)
        path=directory/('coupons-1904.xlsx' if mac else 'coupons-1900.xlsx');book.save(path)
        with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
        ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}';xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for cell in xml.iter(ns+'c'):
            if cell.find(ns+'f') is not None:
                v=cell.find(ns+'v')
                if v is None:v=ET.SubElement(cell,ns+'v')
                v.text='901'
        parts['xl/worksheets/sheet1.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
            for name,value in parts.items():z.writestr(name,value)
        paths.append(path)
    return paths
def verify(path,test_bin_dir,directory,frequency=2):
    path=Path(path);directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    before=path.read_bytes();saved=directory/(path.stem+'-recalculated.ods')
    out=subprocess.check_output([str(Path(test_bin_dir).resolve()/'calligra-native-financial-test'),'--document',str(path),str(saved)],text=True,
                                env={**os.environ,'LD_LIBRARY_PATH':'/usr/local/lib/domed/bitmap'})
    rows=json.loads(out);incomplete={13,14,15};dates={1,6,8}
    for column,(got,want) in enumerate(zip(rows[0],expected(frequency)),1):
        if column in dates:assert got['display'] in (want,want+'T00:00:00.000Z'),(path,column,got,want)
        elif isinstance(want,str):assert got['display']==want,(path,column,got,want)
        elif isinstance(want,bool):assert got['kind']==2 and got['display']==('TRUE' if want else 'FALSE'),(path,column,got,want)
        else:assert got['kind']==1 and math.isclose(got['value'],want,rel_tol=1e-12,abs_tol=1e-10),(path,column,got,want)
        assert got['incomplete']==(column in incomplete),(column,got)
    for row,date in zip(rows,DATES):assert row[0]['value']==(date-(dt.date(1904,1,1) if path.suffix=='.xlsx' and '1904' in path.stem else dt.date(1899,12,30))).days
    office='{urn:oasis:names:tc:opendocument:xmlns:office:1.0}';table='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    with zipfile.ZipFile(saved) as z:
        root=ET.fromstring(z.read('content.xml'));cells=next(root.iter(table+'table-row')).findall(table+'table-cell')
        for c in incomplete:assert cells[c-1].get(office+'value')=='901',(c,cells[c-1].attrib)
        for c,want in enumerate(expected(frequency),1):
            cell=cells[c-1]
            if c in dates:
                assert cell.get(office+'value-type')=='date' and cell.get(office+'date-value') in (want,want+'T00:00:00.000Z'),(c,cell.attrib)
                assert cell.get(office+'value') is None,(c,cell.attrib)
            elif isinstance(want,bool):assert cell.get(office+'value-type')=='boolean' and cell.get(office+'boolean-value')=='true'
            elif not isinstance(want,str) and c not in incomplete:assert math.isclose(float(cell.get(office+'value')),want,rel_tol=1e-12,abs_tol=1e-10),(c,cell.attrib)
    assert path.read_bytes()==before
    (directory/(path.stem+'-report.json')).write_text(json.dumps(rows,indent=2)+'\n')
    print('PASS',path.name,'coupon schedules, typed coupon dates, epochs, stale cache replacement and unsupported-cache retention',flush=True)
    return expected(frequency)
if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('directory',type=Path);p.add_argument('--test-bin-dir',type=Path,default=ROOT/'build-native/fltk');a=p.parse_args()
    for path in produce(a.directory):verify(path,a.test_bin_dir,a.directory)
