#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""ODF/Excel engineering formulas produced by independent document libraries."""
import argparse,json,os,subprocess
from pathlib import Path
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableCell,TableRow
from odf.text import P
from openpyxl import Workbook
ROOT=Path(__file__).resolve().parents[2]
p=argparse.ArgumentParser(description=__doc__);p.add_argument('directory',type=Path);p.add_argument('--test-bin-dir',type=Path,default=ROOT/'build-native/fltk');a=p.parse_args()
a.directory.mkdir(parents=True,exist_ok=True)
formulas=['COMPLEX(A1;4;"j")','IMABS(B1)','CONVERT(A1;"m";"cm")','DEC2HEX(A1+12)','IMREAL(IMPRODUCT(B1;COMPLEX(1;-1)))','BESSELJ(A1;2)','IMSUM({"1+2i";"3-4i"})','COM.SUN.STAR.SHEET.ADDIN.ANALYSIS.GETIMEXP(0)','IMREAL(IMPOWER(COMPLEX(0;1);2))','ERF(8;9)','CONVERT(1;"MTON";"ft3")','UNKNOWN_ENGINEERING()']
book=OpenDocumentSpreadsheet();sheet=Table(name='External');book.spreadsheet.addElement(sheet);row=TableRow();sheet.addElement(row);row.addElement(TableCell(valuetype='float',value=3))
for i,f in enumerate(formulas,2):
 # Deliberately stale cached values must be replaced except for unknown formulas.
 cell=TableCell(formula='of:='+f,valuetype='float',value=901);cell.addElement(P(text='901'));row.addElement(cell)
ods=a.directory/'external-engineering.ods';book.save(str(ods))
xlsx=Workbook();ws=xlsx.active;ws.title='External';ws['A1']=3
for i,f in enumerate(formulas,2):
 if i==8:f='IMSUM("1+2i";"3-4i")'
 if i==9:f='_xlfn.IMEXP(0)'
 ws.cell(1,i,'='+f.replace(';',','))
xpath=a.directory/'external-engineering.xlsx';xlsx.save(xpath)
for path in [ods,xpath]:
 before=path.read_bytes()
 raw=subprocess.check_output([str(a.test_bin_dir.resolve()/'calligra-native-engineering-test'),'--document',str(path)],text=True,env={**os.environ,'LD_LIBRARY_PATH':'/usr/local/lib/domed/bitmap'})
 row=json.loads(raw)[0]
 assert row[0]['real']==3 and row[1]['kind']=='complex' and row[1]['real']==3 and row[1]['imag']==4 and row[1]['suffix']=='j'
 for at,value in [(2,5),(3,300),(5,7),(9,-1),(11,40)]:assert abs(row[at]['real']-value)<1e-12,(path,at,row[at])
 assert row[4]['kind']=='string' and row[4]['display']=='F'
 assert row[7]['kind']=='complex' and row[7]['real']==4 and row[7]['imag']==-2,row[7]
 assert row[8]['kind']=='complex' and row[8]['real']==1 and row[8]['imag']==0,row[8]
 assert row[10]['real']>1e-30
 assert row[12]['kind']=='error' and row[12]['incomplete']
 assert before==path.read_bytes()
 print('PASS:',path.suffix,'independent engineering formulas, legacy/Excel namespaces, typed complex and unsupported status',flush=True)
