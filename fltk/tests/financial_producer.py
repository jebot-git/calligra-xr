#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-2.0-or-later
"""Independent ODS/XLSX financial documents, recalculation and cache retention."""
import argparse,json,math,os,re,subprocess,zipfile
from pathlib import Path
import xml.etree.ElementTree as ET
import numpy_financial as npf
import numpy as np
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table,TableCell,TableRow
from odf.text import P
from openpyxl import Workbook

ROOT=Path(__file__).resolve().parents[2]
FORMULAS=[
 'PMT(.005;60;-[.A1])', 'RATE(60;[.B1];-[.A1])', 'IRR([.A1:.A3])',
 'MIRR([.A1:.A3];.1;.12)', 'CUMIPMT(.005;60;-[.A1];1;12;0)',
 'CUMPRINC(.005;60;-[.A1];1;12;0)', 'PPMT(.005;10;60;-[.A1])',
 'FVSCHEDULE(1000;{.03;.04;.05})', 'VDB(10000;600;10;4.875;5.875;1.5)',
 'NPER(.005;[.B1];-[.A1])', 'RATE(0;1;2)',
 'ODDLPRICE(DATE(2026;1;1);DATE(2030;1;1);DATE(2025;7;1);.05;.07;100;2)', 'IFERROR([.M1];77)',
 'UNKNOWN_FINANCIAL()', 'PMT(1E-20;10;1000)', 'IFERROR([.L1];77)',
 'RRI(.5;100;200)', 'PV_ANNUITY(1000;.05;5)', 'NPV(.1;[.A1:.A3])', '[.C1]+.001']

def expected(cash=-1000):
    pv=-cash;flows=[cash,600,700];periods=np.arange(1,13)
    return [cash,float(npf.pmt(.005,60,pv)),.005,float(npf.irr(flows)),
            float(npf.mirr(flows,.1,.12)),float(np.sum(npf.ipmt(.005,periods,60,pv))),
            float(np.sum(npf.ppmt(.005,periods,60,pv))),float(npf.ppmt(.005,10,60,pv)),
            1124.76,769.36046875,60,'#VALUE!','#NAME?',77,'#NAME?',-100,77,3,
            float(npf.pv(.05,5,-1000)),float(npf.npv(.1,flows)/1.1),.006]

def produce(directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=OpenDocumentSpreadsheet();sheet=Table(name='Financial');doc.spreadsheet.addElement(sheet)
    for r,cash in enumerate([-1000,600,700]):
        row=TableRow();sheet.addElement(row);row.addElement(TableCell(valuetype='float',value=cash))
        if r==0:
            for formula in FORMULAS:
                cell=TableCell(formula='of:='+formula,valuetype='float',value=901)
                cell.addElement(P(text='901'));row.addElement(cell)
    ods=directory/'external-financial.ods';doc.save(str(ods))
    book=Workbook();ws=book.active;ws.title='Financial'
    for r,cash in enumerate([-1000,600,700],1):ws.cell(r,1,cash)
    for c,formula in enumerate(FORMULAS,2):
        f=re.sub(r'\[\.([A-Z]+[0-9]+)(?::\.([A-Z]+[0-9]+))?\]',lambda m:m[1]+(':'+m[2] if m[2] else ''),formula).replace(';',',')
        if f.startswith('RRI('):f='_xlfn.'+f
        ws.cell(1,c,'='+f)
    xlsx=directory/'external-financial.xlsx';book.save(xlsx)
    # Seed stale numeric caches in the independent producer's XLSX, too.
    with zipfile.ZipFile(xlsx) as source:parts={name:source.read(name) for name in source.namelist()}
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    xml=ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    for cell in xml.iter(ns+'c'):
        if cell.find(ns+'f') is None:continue
        value=cell.find(ns+'v')
        if value is None:value=ET.SubElement(cell,ns+'v')
        value.text='901'
    parts['xl/worksheets/sheet1.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(xlsx,'w',zipfile.ZIP_DEFLATED) as target:
        for name,value in parts.items():target.writestr(name,value)
    return [ods,xlsx]

def verify(path,test_bin_dir,directory,cash=-1000):
    path=Path(path);directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    before=path.read_bytes();saved=directory/(path.suffix[1:]+'-recalculated.ods')
    command=[str(Path(test_bin_dir).resolve()/'calligra-native-financial-test'),'--document',str(path),str(saved)]
    run=subprocess.run(command,text=True,capture_output=True,env={**os.environ,'LD_LIBRARY_PATH':'/usr/local/lib/domed/bitmap'},check=True)
    rows=json.loads(run.stdout);want=expected(cash)
    for c,(got,value) in enumerate(zip(rows[0],want),1):
        if isinstance(value,str):assert got['display']==value,(path,c,got,value)
        else:assert got['kind']==1 and math.isclose(got['value'],value,rel_tol=3e-10,abs_tol=1e-10),(path,c,got,value)
        assert got['incomplete']==(c in (13,14,15)),(path,c,got)
    office='{urn:oasis:names:tc:opendocument:xmlns:office:1.0}';table='{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
    with zipfile.ZipFile(saved) as archive:
        root=ET.fromstring(archive.read('content.xml'));cells=next(root.iter(table+'table-row')).findall(table+'table-cell')
        assert len(cells)>=21
        for c in (13,14,15):assert cells[c-1].get(office+'value')=='901',(path,c,cells[c-1].attrib)
        assert cells[11].get(office+'string-value')=='#VALUE!',cells[11].attrib
        assert cells[16].get(office+'value')=='77',cells[16].attrib
        for c,value in enumerate(want,1):
            if isinstance(value,str) or c in (13,14,15):continue
            assert math.isclose(float(cells[c-1].get(office+'value')),value,rel_tol=3e-10,abs_tol=1e-10),(path,c,cells[c-1].attrib)
    assert before==path.read_bytes()
    (directory/(path.suffix[1:]+'-report.json')).write_text(json.dumps(rows,indent=2)+'\n')
    print('PASS',path.suffix,'independent financial formulas, amortization, IRR, tiny rates, recalculated caches and unsupported coupon preservation',flush=True)
    return want

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('directory',type=Path)
    parser.add_argument('--test-bin-dir',type=Path,default=ROOT/'build-native/fltk');args=parser.parse_args()
    for path in produce(args.directory):verify(path,args.test_bin_dir,args.directory)
